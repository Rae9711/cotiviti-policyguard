#!/usr/bin/env python3
"""Generate the PolicyGuard assessment PowerPoint deck (python-pptx)."""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "slides" / "PolicyGuard_Presentation.pptx"
EVAL_JSON = ROOT / "evaluation" / "results.json"

INK = RGBColor(0x1A, 0x23, 0x32)
TEAL = RGBColor(0x0D, 0x6E, 0x6E)
MUTED = RGBColor(0x5C, 0x6B, 0x7A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xF4, 0xF6, 0xF8)


def _pct(v: float | None) -> str:
    if v is None:
        return "n/a"
    return f"{100 * v:.0f}%"


def _set_run(run, text: str, size: int, bold: bool = False, color=INK) -> None:
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def _fix_title(slide, prs, title, subtitle=None):
    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(0.85))
    bar.fill.solid()
    bar.fill.fore_color.rgb = INK
    bar.line.fill.background()
    box = slide.shapes.add_textbox(Inches(0.45), Inches(0.2), Inches(9.2), Inches(0.5))
    tf = box.text_frame
    run = tf.paragraphs[0].add_run()
    _set_run(run, title, 22, True, WHITE)
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.45), Inches(1.0), Inches(9.2), Inches(0.35))
        r = sub.text_frame.paragraphs[0].add_run()
        _set_run(r, subtitle, 12, False, MUTED)


def _bullets(slide, left, top, width, height, items, size=16):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(size)
        p.font.color.rgb = INK
        p.font.name = "Calibri"
        p.space_after = Pt(8)


def build() -> None:
    metrics = {}
    if EVAL_JSON.exists():
        metrics = json.loads(EVAL_JSON.read_text(encoding="utf-8"))
    cd = metrics.get("change_detection", {})
    ab = metrics.get("abstention", {})
    rt = metrics.get("rule_tests", {})

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 Cover
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = INK
    bg.line.fill.background()
    accent = s.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.18), prs.slide_height)
    accent.fill.solid()
    accent.fill.fore_color.rgb = TEAL
    accent.line.fill.background()
    t = s.shapes.add_textbox(Inches(0.7), Inches(2.2), Inches(8.5), Inches(1))
    r = t.text_frame.paragraphs[0].add_run()
    _set_run(r, "POLICYGUARD", 36, True, WHITE)
    st = s.shapes.add_textbox(Inches(0.7), Inches(3.1), Inches(8.5), Inches(1.2))
    p = st.text_frame.paragraphs[0]
    r = p.add_run()
    _set_run(
        r,
        "Source-grounded policy change detection and declarative rule validation\n"
        "for healthcare payment content management",
        16,
        False,
        RGBColor(0xC8, 0xD2, 0xDC),
    )
    foot = s.shapes.add_textbox(Inches(0.7), Inches(6.4), Inches(8.5), Inches(0.6))
    fr = foot.text_frame.paragraphs[0].add_run()
    _set_run(
        fr,
        "Cotiviti Generative AI Research Intern Assessment  ·  Topic 3  ·  Haorui Wang",
        12,
        False,
        RGBColor(0x9A, 0xA8, 0xB5),
    )

    # 2 Problem
    s = prs.slides.add_slide(blank)
    _fix_title(s, prs, "Problem", "Payment policy changes must become validated, reviewable rules")
    _bullets(
        s,
        0.6,
        1.5,
        8.8,
        5,
        [
            "Billing/coding policies change continuously (e.g., CMS NCCI-style guidance).",
            "Material edits include effective dates and qualitative “integral service” language.",
            "Risk A: missed updates → payment leakage. Risk B: unsupervised automation → weak accountability.",
            "Need: evidence-grounded change detection + constrained rules + human approval.",
        ],
    )

    # 3 Approach
    s = prs.slides.add_slide(blank)
    _fix_title(s, prs, "Approach", "Deterministic first; LLM optional and non-required")
    _bullets(
        s,
        0.6,
        1.5,
        8.8,
        5,
        [
            "difflib policy version comparison (offline, auditable).",
            "Source-grounded change objects: document, section, effective date, evidence.",
            "Pydantic JSON rules — not executable Python; actions limited to flag_for_review.",
            "Ambiguous language → abstain / request expert interpretation.",
            "Append-only JSONL audit of reviewer decisions.",
        ],
    )

    # 4 Architecture
    s = prs.slides.add_slide(blank)
    _fix_title(s, prs, "Architecture", "PolicyGuard pipeline")
    _bullets(
        s,
        0.6,
        1.5,
        8.8,
        5,
        [
            "app.py (Streamlit) orchestrates select → diff → analyze → review → impact → audit.",
            "src/policy_diff.py → src/change_analyzer.py → src/schema.py → src/rule_engine.py.",
            "src/audit_log.py persists HITL decisions; evaluation/ scores a demo gold set.",
            "Demo runs without an API key using cached/deterministic rule proposals.",
        ],
    )

    # 5 Demo walkthrough
    s = prs.slides.add_slide(blank)
    _fix_title(s, prs, "Demo walkthrough", "Synthetic NCCI-style 2025 → 2026 excerpts")
    _bullets(
        s,
        0.6,
        1.5,
        8.8,
        5,
        [
            "EXAMPLE1: machine-clear effective-date revision → proposed declarative rule.",
            "Approve → deterministic engine flags potentially affected synthetic claims.",
            "EXAMPLE2 integral language: abstain — insufficient evidence for automation.",
            "Careful language only: flagged for review; no denial / fraud claims.",
        ],
    )

    # 6 Evaluation
    s = prs.slides.add_slide(blank)
    _fix_title(
        s,
        prs,
        "Evaluation (POC-scale)",
        "Actual metrics from evaluation/results — tiny gold set, honest scope",
    )
    _bullets(
        s,
        0.6,
        1.5,
        8.8,
        5,
        [
            f"Change-detection precision/recall: {_pct(cd.get('precision'))} / {_pct(cd.get('recall'))}",
            f"Abstention correctness: {_pct(ab.get('accuracy'))}",
            f"Rule-test pass rate: {_pct(rt.get('pass_rate'))}",
            "Also tracked: effective-date accuracy, source-citation coverage, simulated reviewer rates.",
            "Not a production performance claim — instrumentation for assessment demo.",
        ],
    )

    # 7 Governance
    s = prs.slides.add_slide(blank)
    _fix_title(s, prs, "Governance & HITL", "Expert accountability over autonomous payment action")
    _bullets(
        s,
        0.6,
        1.5,
        8.8,
        5,
        [
            "Reviewer gates: Approve / Reject / Request Expert Interpretation.",
            "Schema forbids denial-style actions; engine cannot auto-deny claims.",
            "Abstention is a success mode for ambiguous policy language.",
            "Audit JSONL supports reconstruction of who approved what, when.",
        ],
    )

    # 8 Recommendation
    s = prs.slides.add_slide(blank)
    _fix_title(s, prs, "Recommendation", "Strategic fit for payment policy management")
    _bullets(
        s,
        0.6,
        1.5,
        8.8,
        5,
        [
            "Near term: deterministic diff + constrained rules + HITL queue for clear edits.",
            "Mid term: optional LLM proposals with retrieval + schema validation + abstention.",
            "Do not execute model-authored Python in payment paths.",
            "Measure abstention quality and reviewer acceptance alongside detection metrics.",
        ],
    )

    # 9 Closing
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = INK
    bg.line.fill.background()
    t = s.shapes.add_textbox(Inches(0.7), Inches(2.5), Inches(8.5), Inches(1.5))
    r = t.text_frame.paragraphs[0].add_run()
    _set_run(r, "Policy → validated rule → expert accountability", 24, True, WHITE)
    st = s.shapes.add_textbox(Inches(0.7), Inches(4.0), Inches(8.5), Inches(1.2))
    p = st.text_frame.paragraphs[0]
    r = p.add_run()
    _set_run(
        r,
        "Disclaimer: PolicyGuard is a HITL proof of concept. It does not make autonomous\n"
        "clinical, payment, or claim-denial decisions.",
        13,
        False,
        RGBColor(0xC8, 0xD2, 0xDC),
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
